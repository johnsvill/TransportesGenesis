"""Genera la versión Word de la tesis desde TESIS_TRANSPORTES_GENESIS.md.

El formato sigue la guía institucional proporcionada: tamaño carta, Times New
Roman 12, doble espacio, márgenes superior/izquierdo de 4 cm y
inferior/derecho de 2.5 cm, preliminares romanos y cuerpo arábigo.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


BASE = Path(__file__).resolve().parent
SOURCE = BASE / "TESIS_TRANSPORTES_GENESIS.md"
OUTPUT = BASE / "TESIS_TRANSPORTES_GENESIS.docx"

FONT = "Times New Roman"
FONT_SIZE = Pt(12)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline(paragraph, text, force_bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, *, bold=None, italic=None, size=FONT_SIZE) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text: str, *, force_bold: bool = False) -> None:
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, bold=True if force_bold else None)


def set_field(run, instruction: str, placeholder: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_number(paragraph, alignment) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    set_run_font(run, size=Pt(10))
    set_field(run, " PAGE ", "1")


def set_page_number_format(section, fmt: str, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), fmt)
    pg_num.set(qn("w:start"), str(start))


def unlink_headers_and_footers(section) -> None:
    for part in (
        section.header,
        section.first_page_header,
        section.even_page_header,
        section.footer,
        section.first_page_footer,
        section.even_page_footer,
    ):
        part.is_linked_to_previous = False


def set_margins(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(4)
    section.left_margin = Cm(4)
    section.bottom_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def configure_cover_section(section) -> None:
    set_margins(section)
    unlink_headers_and_footers(section)
    section.different_first_page_header_footer = True


def configure_preliminary_section(section) -> None:
    set_margins(section)
    unlink_headers_and_footers(section)
    set_page_number_format(section, "lowerRoman", 2)
    add_page_number(section.footer.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)


def configure_body_section(section) -> None:
    set_margins(section)
    unlink_headers_and_footers(section)
    section.different_first_page_header_footer = True
    set_page_number_format(section, "decimal", 1)
    add_page_number(section.header.paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT)
    add_page_number(
        section.even_page_header.paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT
    )


def configure_continued_body_section(section) -> None:
    set_margins(section)
    section.different_first_page_header_footer = False


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    even_odd = settings.find(qn("w:evenAndOddHeaders"))
    if even_odd is None:
        even_odd = OxmlElement("w:evenAndOddHeaders")
        settings.append(even_odd)
    even_odd.set(qn("w:val"), "true")


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = FONT_SIZE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.first_line_indent = Cm(0.63)
    normal.paragraph_format.space_after = Pt(0)

    for name, size in (("Heading 1", Pt(14)), ("Heading 2", Pt(12)), ("Heading 3", Pt(12))):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = size
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 3"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "Tesis Caption" not in styles:
        caption = styles.add_style("Tesis Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Tesis Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption.font.size = Pt(10)
    caption.font.bold = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = FONT_SIZE
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        style.paragraph_format.space_after = Pt(0)


def parse_frontmatter(raw: str) -> tuple[dict[str, str], str]:
    if not raw.startswith("---"):
        return {}, raw
    end = raw.find("\n---", 3)
    if end < 0:
        return {}, raw
    meta = {}
    for line in raw[3:end].strip().splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip()] = value.strip().strip('"')
    return meta, raw[end + 4 :].lstrip()


def add_hidden_tc(paragraph, text: str, category: str) -> None:
    run = paragraph.add_run()
    escaped = text.replace('"', "'")
    run.font.hidden = True
    set_field(run, f'TC "{escaped}" \\f {category} \\l 1')


def add_caption(doc: Document, text: str) -> None:
    clean = re.sub(r"\*\*", "", text).strip()
    paragraph = doc.add_paragraph(style="Tesis Caption")
    add_inline(paragraph, clean, force_bold=True)
    category = "C" if clean.lower().startswith("cuadro") else "F"
    add_hidden_tc(paragraph, clean, category)


def add_toc_field(doc: Document, category: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run()
    set_run_font(run)
    instruction = "TOC \\h \\z \\u"
    if category:
        instruction = f"TOC \\h \\z \\f {category}"
    set_field(run, instruction, "Actualice este campo en Word.")


def add_heading(doc: Document, text: str, level: int) -> None:
    clean = re.sub(r"\*\*", "", text).strip()
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        clean = clean.upper()
    run = paragraph.add_run(clean)
    set_run_font(run, bold=True, size=Pt(14) if level == 1 else FONT_SIZE)


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    add_inline(paragraph, text.strip())


def add_list_item(doc: Document, text: str, numbered: bool) -> None:
    style = "List Number" if numbered else "List Bullet"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_inline(paragraph, text)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            value = row[col_index] if col_index < len(row) else ""
            set_cell_text(
                table.cell(row_index, col_index),
                re.sub(r"\*\*", "", value),
                bold=row_index == 0,
            )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.first_line_indent = Cm(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def build_document() -> Document:
    raw = SOURCE.read_text(encoding="utf-8")
    _, body = parse_frontmatter(raw)
    doc = Document()
    configure_styles(doc)
    configure_cover_section(doc.sections[0])
    set_update_fields(doc)

    lines = body.splitlines()
    index = 0
    table_rows: list[list[str]] = []
    code_lines: list[str] = []
    in_code = False
    cover_done = False
    body_started = False
    first_body_unit = True

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if in_code:
            if stripped.startswith("```"):
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("```"):
            flush_table()
            in_code = True
            index += 1
            continue

        if stripped.startswith("|") and not re.match(r"^\|\s*[-:| ]+\|\s*$", stripped):
            table_rows.append([cell.strip() for cell in stripped.strip("|").split("|")])
            index += 1
            if index < len(lines) and re.match(r"^\|\s*[-:| ]+\|\s*$", lines[index].strip()):
                index += 1
            continue
        if table_rows:
            flush_table()

        if stripped == "<!-- PAGE_BREAK -->":
            if not cover_done:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_preliminary_section(section)
                cover_done = True
            else:
                doc.add_page_break()
            index += 1
            continue

        if stripped == "<!-- BODY_START -->":
            section = doc.add_section(WD_SECTION.ODD_PAGE)
            configure_body_section(section)
            body_started = True
            index += 1
            continue

        if not stripped or stripped == "---":
            index += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            begins_odd_page = (
                body_started
                and level == 1
                and (
                    text.upper().startswith("CAPÍTULO")
                    or text.upper().startswith("APÉNDICE")
                )
            )
            if begins_odd_page:
                if first_body_unit:
                    first_body_unit = False
                else:
                    section = doc.add_section(WD_SECTION.ODD_PAGE)
                    configure_continued_body_section(section)
            add_heading(doc, text, level)
            upper = re.sub(r"\*\*", "", text).upper()
            if upper == "ÍNDICE GENERAL":
                add_toc_field(doc)
            elif upper == "ÍNDICE DE CUADROS":
                add_toc_field(doc, "C")
            elif upper == "ÍNDICE DE FIGURAS":
                add_toc_field(doc, "F")
            index += 1
            continue

        if re.match(r"^\*\*(Cuadro|Figura)\s+\d+\.\d+\.", stripped, re.IGNORECASE):
            add_caption(doc, stripped)
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            add_list_item(doc, numbered.group(1), numbered=True)
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            add_list_item(doc, bullet.group(1), numbered=False)
            index += 1
            continue

        add_body_paragraph(doc, stripped)
        index += 1

    flush_table()
    if in_code:
        add_code_block(doc, code_lines)
    return doc


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"No existe el archivo fuente: {SOURCE}")
    document = build_document()
    document.save(OUTPUT)
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
