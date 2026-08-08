"""Genera la presentación y el guion oral de la tesis Transportes Génesis."""

from __future__ import annotations

import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parent
OUT = BASE / "PRESENTACION_DEFENSA_TESIS_TRANSPORTES_GENESIS.pptx"
PDF_OUT = BASE / "PRESENTACION_DEFENSA_TESIS_TRANSPORTES_GENESIS.pdf"
SCRIPT_OUT = BASE / "GUION_PRESENTACION_TESIS.md"
DIAGRAMS = BASE / "diagramas_integrados"

ARCH = DIAGRAMS / "figura_4_2_arquitectura_general.png"
ATTENDANCE = DIAGRAMS / "figura_4_5_asistencia_abordaje.png"

NAVY = RGBColor(10, 30, 52)
BLUE = RGBColor(30, 91, 168)
CYAN = RGBColor(26, 166, 183)
AMBER = RGBColor(238, 161, 61)
GREEN = RGBColor(38, 139, 87)
YELLOW = RGBColor(230, 169, 45)
RED = RGBColor(184, 67, 67)
INK = RGBColor(32, 45, 61)
GRAY = RGBColor(98, 112, 127)
LIGHT = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(215, 224, 233)

SLIDE_W = Inches(13.333333)
SLIDE_H = Inches(7.5)

SLIDES = [
    (
        "Portada",
        "Presentamos la implementación y validación funcional de una plataforma web "
        "para Transportes Génesis. El trabajo integra rutas, geolocalización y "
        "confirmación diaria de asistencia. La defensa diferencia lo observado de "
        "lo que todavía requiere pruebas adicionales.",
    ),
    (
        "Problema y contexto",
        "La operación coordina buses, rutas, paradas, estudiantes, pilotos, monitores "
        "y familias. Cuando la información se reparte entre llamadas, mensajes y "
        "registros separados, se dificulta consultar la ubicación, la ruta y el estado "
        "diario del estudiante. La propuesta centraliza esas funciones por roles.",
    ),
    (
        "Hipótesis y objetivos",
        "La hipótesis se limita al funcionamiento observable. No afirma reducción de "
        "tiempos, satisfacción, seguridad física ni adopción productiva. Buscamos "
        "verificar si rutas, ubicación y asistencia pueden coexistir y consultarse en "
        "una plataforma común.",
    ),
    (
        "Solución y arquitectura",
        "La solución se organiza por capas. Las interfaces se conectan con páginas, "
        "controladores y API; los servicios concentran la lógica; Entity Framework "
        "Core accede a SQL Server; y SignalR distribuye actualizaciones a los clientes. "
        "Leaflet, OpenStreetMap y OSRM apoyan la representación cartográfica.",
    ),
    (
        "Roles y flujo operativo",
        "El administrador configura la operación; piloto y monitor consultan el "
        "recorrido; y el padre accede al mapa y confirma asistencia por fecha y turno. "
        "La optimización de paradas usa una heurística práctica, por lo que no se "
        "presenta como una solución globalmente óptima.",
    ),
    (
        "Geolocalización en tiempo casi real",
        "El navegador obtiene la ubicación del piloto o monitor. La API recibe las "
        "coordenadas, el servicio puede persistirlas y SignalR distribuye la "
        "actualización al mapa. Usamos el término tiempo casi real porque existe "
        "latencia y la precisión depende del dispositivo, permisos y conectividad.",
    ),
    (
        "Asistencia y abordaje",
        "La asistencia representa la intención de utilizar el transporte en una fecha "
        "y turno; el abordaje representa una recogida operativa. Los videos respaldan "
        "la interfaz de asistencia del padre. No se recibió evidencia inequívoca del "
        "monitor guardando y recuperando un registro de abordaje.",
    ),
    (
        "Metodología y evidencia",
        "La investigación es aplicada, descriptiva y evaluativa bajo un estudio de "
        "caso tecnológico. Se revisaron documentos, dos casos funcionales, evidencia "
        "audiovisual y arquitectura. No se realizaron encuestas ni se derivaron "
        "porcentajes de satisfacción.",
    ),
    (
        "Pruebas y escenarios",
        "El informe del 4 de agosto documenta dos casos del administrador y declara "
        "ambos aprobados. Sin embargo, una captura no muestra el mapa esperado y la "
        "otra no demuestra una edición guardada. Por eso distinguimos el estado "
        "declarado de la verificación visual independiente.",
    ),
    (
        "Resultados y alcance real",
        "El resultado comprobable es la integración observable de rutas, mapas, "
        "ubicación y asistencia diaria en un ambiente controlado. El abordaje "
        "persistente, el rendimiento, la producción y la aceptación integral no "
        "fueron demostrados.",
    ),
    (
        "Evaluación de hipótesis y trabajo futuro",
        "La hipótesis queda respaldada dentro del alcance funcional documentado, no "
        "como certificación integral. El siguiente paso es verificar persistencia, "
        "ampliar pruebas con padres, pilotos y monitores, evaluar autorizaciones "
        "negativas y medir latencia, carga y estabilidad.",
    ),
    (
        "Cierre",
        "El aporte verificable es una plataforma integrada y una evidencia funcional "
        "trazable, acompañada de límites explícitos. Fueron observables rutas, mapas, "
        "geolocalización y la interfaz de asistencia; abordaje y pagos requieren "
        "validación adicional. Muchas gracias.",
    ),
]


def set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, kind, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(kind, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=20,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    font="Aptos",
    valign=MSO_ANCHOR.TOP,
    margin=0.06,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=18, color=INK, gap=9):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.12)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.05)
    for index, line in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.level = 0
    return box


def add_header(slide, number: int, title: str, kicker: str) -> None:
    add_text(slide, f"{number:02d}", Inches(0.55), Inches(0.33), Inches(0.65), Inches(0.35), 13, CYAN, True)
    add_text(slide, kicker.upper(), Inches(1.15), Inches(0.33), Inches(4.5), Inches(0.35), 11, GRAY, True)
    add_text(slide, title, Inches(0.55), Inches(0.82), Inches(12.1), Inches(0.65), 28, NAVY, True)
    rule = add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.53), Inches(12.2), Inches(0.025), CYAN)
    rule.line.fill.background()


def add_footer(slide, number: int, source: str = "Fuente: elaboración propia con base en la tesis final.") -> None:
    add_text(slide, source, Inches(0.55), Inches(7.13), Inches(10.9), Inches(0.22), 8, GRAY)
    add_text(slide, str(number), Inches(12.15), Inches(7.08), Inches(0.55), Inches(0.28), 9, GRAY, True, PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, accent=BLUE, title_size=16, body_size=13):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, LINE)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(0.08), h, accent, accent)
    add_text(slide, title, x + Inches(0.25), y + Inches(0.18), w - Inches(0.45), Inches(0.36), title_size, NAVY, True)
    add_text(slide, body, x + Inches(0.25), y + Inches(0.64), w - Inches(0.45), h - Inches(0.78), body_size, INK)


def add_tag(slide, text, x, y, w, color=BLUE, light=LIGHT):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.38), light, light)
    add_text(slide, text, x, y + Inches(0.02), w, Inches(0.28), 10, color, True, PP_ALIGN.CENTER)


def add_notes(slide, text: str) -> None:
    try:
        frame = slide.notes_slide.notes_text_frame
        frame.text = text
    except (AttributeError, ValueError):
        pass


def slide_portada(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), SLIDE_H, CYAN)
    add_text(slide, "TESIS · DEFENSA FINAL", Inches(0.8), Inches(0.65), Inches(4.2), Inches(0.4), 13, CYAN, True)
    add_text(
        slide,
        "Transportes\nGénesis",
        Inches(0.8),
        Inches(1.35),
        Inches(5.5),
        Inches(1.7),
        40,
        WHITE,
        True,
    )
    add_text(
        slide,
        "Implementación y validación funcional de una plataforma web de "
        "geolocalización, asistencia y abordaje",
        Inches(0.8),
        Inches(3.18),
        Inches(6.15),
        Inches(1.15),
        21,
        RGBColor(219, 230, 240),
    )
    add_text(
        slide,
        "Jonathan Samuel Villeda Pérez\nJosé David Florián Secaida",
        Inches(0.8),
        Inches(5.08),
        Inches(5.7),
        Inches(0.82),
        16,
        WHITE,
        True,
    )
    add_text(slide, "Universidad Galileo · Ingeniería en Ciencias y Sistemas", Inches(0.8), Inches(6.06), Inches(6.4), Inches(0.35), 11, RGBColor(182, 201, 217))
    add_text(slide, "[Facultad/Escuela] · [Asesor] · Agosto 2026", Inches(0.8), Inches(6.48), Inches(6.4), Inches(0.3), 10, AMBER, True)

    # Visual abstract: bus route moving through an integrated platform.
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.65), Inches(1.0), Inches(4.75), Inches(5.45), RGBColor(19, 47, 76), RGBColor(46, 78, 106))
    add_text(slide, "PLATAFORMA INTEGRADA", Inches(8.02), Inches(1.35), Inches(4.0), Inches(0.4), 12, CYAN, True, PP_ALIGN.CENTER)
    for i, (label, color) in enumerate((("RUTAS", BLUE), ("UBICACIÓN", CYAN), ("ASISTENCIA", AMBER))):
        cy = 2.15 + i * 1.2
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(8.1), Inches(cy), Inches(0.58), Inches(0.58), color, color)
        add_text(slide, str(i + 1), Inches(8.1), Inches(cy + 0.1), Inches(0.58), Inches(0.28), 13, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, label, Inches(8.95), Inches(cy + 0.08), Inches(2.6), Inches(0.35), 17, WHITE, True)
        if i < 2:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8.37), Inches(cy + 0.6), Inches(0.05), Inches(0.62), RGBColor(66, 95, 120))
    add_text(slide, "Alcance: validación funcional documentada", Inches(8.0), Inches(5.9), Inches(4.05), Inches(0.3), 10, RGBColor(182, 201, 217), False, PP_ALIGN.CENTER)
    add_notes(slide, SLIDES[0][1])


def slide_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 2, "Problema y contexto", "Punto de partida")
    add_text(slide, "Información operativa dispersa", Inches(0.72), Inches(1.82), Inches(4.6), Inches(0.45), 19, NAVY, True)
    items = [
        ("Llamadas", "Ubicación comunicada de forma aislada"),
        ("Mensajes", "Cambios de ruta sin trazabilidad común"),
        ("Registros", "Asistencia y recogidas separadas"),
    ]
    for i, (title, body) in enumerate(items):
        add_card(slide, Inches(0.72), Inches(2.42 + i * 1.25), Inches(4.9), Inches(0.95), title, body, [BLUE, AMBER, GRAY][i], 15, 12)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(5.92), Inches(3.05), Inches(1.05), Inches(1.4), CYAN, CYAN)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.25), Inches(2.0), Inches(5.05), Inches(3.95), NAVY, NAVY)
    add_text(slide, "PLATAFORMA COMÚN", Inches(7.65), Inches(2.43), Inches(4.25), Inches(0.48), 20, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "Rutas  ·  Mapas  ·  Asistencia", Inches(7.65), Inches(3.12), Inches(4.25), Inches(0.4), 17, CYAN, True, PP_ALIGN.CENTER)
    add_text(
        slide,
        "Objetivo operativo\nCentralizar la consulta y mejorar la trazabilidad por roles.",
        Inches(7.75),
        Inches(3.92),
        Inches(4.05),
        Inches(1.15),
        15,
        WHITE,
        False,
        PP_ALIGN.CENTER,
    )
    add_text(slide, "No equivale a garantizar seguridad física.", Inches(7.65), Inches(5.36), Inches(4.25), Inches(0.3), 10, AMBER, True, PP_ALIGN.CENTER)
    add_footer(slide, 2)
    add_notes(slide, SLIDES[1][1])


def slide_hypothesis(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 3, "Hipótesis y objetivos", "Alcance verificable")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(1.9), Inches(7.7), Inches(1.55), NAVY, NAVY)
    add_text(slide, "HIPÓTESIS DE INVESTIGACIÓN", Inches(1.02), Inches(2.13), Inches(3.6), Inches(0.3), 11, CYAN, True)
    add_text(
        slide,
        "La integración funcional de geolocalización y asistencia permite "
        "centralizar el monitoreo de buses y rutas y representar confirmaciones diarias.",
        Inches(1.02),
        Inches(2.55),
        Inches(6.95),
        Inches(0.64),
        17,
        WHITE,
        True,
    )
    add_card(slide, Inches(8.72), Inches(1.9), Inches(3.65), Inches(1.55), "Criterio", "Evidencia observable y trazable en ambiente controlado.", CYAN, 15, 13)
    objectives = [
        ("01", "Geolocalización", "Monitoreo de buses y rutas en tiempo casi real."),
        ("02", "Asistencia / abordaje", "Confirmación diaria; abordaje sujeto a evidencia."),
        ("03", "Documentación", "Casos funcionales, material audiovisual y límites."),
    ]
    for i, (num, title, body) in enumerate(objectives):
        x = 0.72 + i * 4.0
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.0), Inches(3.65), Inches(1.62), LIGHT, LINE)
        add_text(slide, num, Inches(x + 0.25), Inches(4.23), Inches(0.55), Inches(0.35), 15, CYAN, True)
        add_text(slide, title, Inches(x + 0.82), Inches(4.2), Inches(2.5), Inches(0.36), 15, NAVY, True)
        add_text(slide, body, Inches(x + 0.25), Inches(4.75), Inches(3.1), Inches(0.55), 12, INK)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.12), Inches(11.65), Inches(0.52), RGBColor(255, 248, 232), RGBColor(248, 222, 165))
    add_text(slide, "NO AFIRMA: reducción de tiempos · satisfacción · seguridad física · producción", Inches(0.95), Inches(6.24), Inches(11.2), Inches(0.25), 11, RGBColor(144, 92, 16), True, PP_ALIGN.CENTER)
    add_footer(slide, 3)
    add_notes(slide, SLIDES[2][1])


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 4, "Solución y arquitectura", "Diseño técnico")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(1.85), Inches(11.92), Inches(4.65), WHITE, LINE)
    slide.shapes.add_picture(str(ARCH), Inches(0.95), Inches(2.15), width=Inches(11.45))
    tags = [
        ("ASP.NET Core 8", Inches(0.85), Inches(6.6), Inches(1.9), BLUE),
        ("SQL Server", Inches(2.92), Inches(6.6), Inches(1.5), CYAN),
        ("SignalR", Inches(4.59), Inches(6.6), Inches(1.25), AMBER),
        ("Leaflet + OSM", Inches(6.01), Inches(6.6), Inches(1.65), GREEN),
        ("EF Core", Inches(7.83), Inches(6.6), Inches(1.25), BLUE),
        ("Identity", Inches(9.25), Inches(6.6), Inches(1.25), GRAY),
    ]
    for text, x, y, w, color in tags:
        add_tag(slide, text, x, y, w, color, LIGHT)
    add_footer(slide, 4, "Fuente: Figura 4.2 de la tesis; diagrama del proyecto.")
    add_notes(slide, SLIDES[3][1])


def slide_roles(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 5, "Roles y flujo operativo", "Responsabilidades")
    roles = [
        ("ADMIN", "Configura buses,\nrutas y paradas", BLUE),
        ("PILOTO", "Consulta ruta y\nreporta ubicación", CYAN),
        ("MONITOR", "Revisa alumnos y\nregistra recogidas", AMBER),
        ("PADRE", "Consulta mapa y\nconfirma asistencia", GREEN),
    ]
    for i, (title, body, color) in enumerate(roles):
        x = 0.72 + i * 3.05
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.95), Inches(2.65), Inches(2.05), WHITE, LINE)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.9), Inches(2.16), Inches(0.85), Inches(0.85), color, color)
        add_text(slide, title[0], Inches(x + 0.9), Inches(2.35), Inches(0.85), Inches(0.3), 18, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, Inches(x + 0.25), Inches(3.18), Inches(2.15), Inches(0.32), 14, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x + 0.25), Inches(3.55), Inches(2.15), Inches(0.5), 11, GRAY, False, PP_ALIGN.CENTER)
    add_text(slide, "Flujo común", Inches(0.72), Inches(4.55), Inches(2.0), Inches(0.35), 16, NAVY, True)
    flow = [("Configurar", BLUE), ("Recorrer", CYAN), ("Reportar", AMBER), ("Consultar", GREEN)]
    for i, (label, color) in enumerate(flow):
        x = 0.85 + i * 3.05
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.13), Inches(2.25), Inches(0.75), color, color)
        add_text(slide, label, Inches(x), Inches(5.34), Inches(2.25), Inches(0.28), 14, WHITE, True, PP_ALIGN.CENTER)
        if i < 3:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x + 2.35), Inches(5.27), Inches(0.5), Inches(0.48), LIGHT, LIGHT)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(6.25), Inches(11.25), Inches(0.45), LIGHT, LIGHT)
    add_text(slide, "Rutas por bus y turno · Orden de paradas mediante heurística de vecino más cercano", Inches(1.0), Inches(6.34), Inches(10.95), Inches(0.25), 11, GRAY, True, PP_ALIGN.CENTER)
    add_footer(slide, 5)
    add_notes(slide, SLIDES[4][1])


def slide_geo(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 6, "Geolocalización en tiempo casi real", "Flujo funcional")
    steps = [
        ("1", "NAVEGADOR", "GPS o simulación", BLUE),
        ("2", "API", "Recibe coordenadas", CYAN),
        ("3", "SERVICIO", "Persiste y emite", AMBER),
        ("4", "MAPA", "Actualiza clientes", GREEN),
    ]
    for i, (num, title, body, color) in enumerate(steps):
        x = 0.68 + i * 3.14
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.15), Inches(2.52), Inches(2.1), WHITE, LINE)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.82), Inches(2.42), Inches(0.88), Inches(0.88), color, color)
        add_text(slide, num, Inches(x + 0.82), Inches(2.63), Inches(0.88), Inches(0.3), 17, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, Inches(x + 0.2), Inches(3.47), Inches(2.12), Inches(0.3), 13, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x + 0.2), Inches(3.82), Inches(2.12), Inches(0.3), 11, GRAY, False, PP_ALIGN.CENTER)
        if i < 3:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x + 2.57), Inches(2.93), Inches(0.44), Inches(0.56), LIGHT, LIGHT)
    add_card(slide, Inches(0.72), Inches(4.86), Inches(3.55), Inches(1.25), "Tecnologías", "watchPosition · API REST · SQL Server · SignalR · Leaflet", BLUE, 15, 12)
    add_card(slide, Inches(4.57), Inches(4.86), Inches(3.55), Inches(1.25), "Interpretación", "“Casi real” reconoce latencia entre captura y visualización.", CYAN, 15, 12)
    add_card(slide, Inches(8.42), Inches(4.86), Inches(3.55), Inches(1.25), "Dependencias", "Permiso, precisión del dispositivo y conectividad.", AMBER, 15, 12)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.37), Inches(11.25), Inches(0.45), RGBColor(245, 249, 252), LINE)
    add_text(slide, "Resultado observado: pantallas de ruta, mapa y ubicación en evidencia audiovisual.", Inches(0.95), Inches(6.46), Inches(10.8), Inches(0.25), 11, GREEN, True, PP_ALIGN.CENTER)
    add_footer(slide, 6)
    add_notes(slide, SLIDES[5][1])


def slide_attendance(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 7, "Asistencia y abordaje", "Dos eventos distintos")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(1.9), Inches(8.15), Inches(4.65), WHITE, LINE)
    slide.shapes.add_picture(str(ATTENDANCE), Inches(0.92), Inches(2.28), width=Inches(7.75))
    add_card(slide, Inches(9.15), Inches(1.9), Inches(3.2), Inches(1.55), "ASISTENCIA", "Confirmación por fecha y turno observable en el perfil del padre.", GREEN, 15, 12)
    add_card(slide, Inches(9.15), Inches(3.72), Inches(3.2), Inches(1.55), "ABORDAJE", "Flujo diseñado; no se demostró una operación guardada y recuperada.", YELLOW, 15, 12)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(9.15), Inches(5.55), Inches(3.2), Inches(1.0), RGBColor(255, 248, 232), RGBColor(248, 222, 165))
    add_text(slide, "No confundir confirmación previa con recogida real.", Inches(9.42), Inches(5.8), Inches(2.65), Inches(0.45), 11, RGBColor(144, 92, 16), True, PP_ALIGN.CENTER)
    add_footer(slide, 7, "Fuente: Figura 4.5 de la tesis. Representa el flujo diseñado.")
    add_notes(slide, SLIDES[6][1])


def slide_method(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 8, "Metodología y evidencia", "Estudio de caso tecnológico")
    labels = [
        ("1", "INVENTARIO", "Reunir documentos y archivos"),
        ("2", "CLASIFICACIÓN", "Ordenar por rol y módulo"),
        ("3", "CONTRASTE", "Comparar esperado y observado"),
        ("4", "LÍMITES", "Separar evidencia y ausencia"),
        ("5", "CONCLUSIÓN", "Responder con alcance explícito"),
    ]
    for i, (num, title, body) in enumerate(labels):
        x = 0.65 + i * 2.48
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.62), Inches(2.0), Inches(0.72), Inches(0.72), [BLUE, CYAN, AMBER, GRAY, GREEN][i], [BLUE, CYAN, AMBER, GRAY, GREEN][i])
        add_text(slide, num, Inches(x + 0.62), Inches(2.17), Inches(0.72), Inches(0.27), 14, WHITE, True, PP_ALIGN.CENTER)
        if i < 4:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x + 1.62), Inches(2.18), Inches(0.48), Inches(0.38), LIGHT, LIGHT)
        add_text(slide, title, Inches(x), Inches(2.98), Inches(1.95), Inches(0.32), 11, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x), Inches(3.4), Inches(1.95), Inches(0.65), 11, GRAY, False, PP_ALIGN.CENTER)
    add_card(slide, Inches(0.72), Inches(4.55), Inches(3.55), Inches(1.42), "Enfoque", "Aplicado · descriptivo · evaluativo", BLUE, 15, 13)
    add_card(slide, Inches(4.57), Inches(4.55), Inches(3.55), Inches(1.42), "Fuentes", "Informe funcional · videos · fotografías · arquitectura", CYAN, 15, 13)
    add_card(slide, Inches(8.42), Inches(4.55), Inches(3.55), Inches(1.42), "No disponible", "Encuestas · carga · acta formal · uso productivo", AMBER, 15, 13)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.28), Inches(11.25), Inches(0.45), LIGHT, LIGHT)
    add_text(slide, "Regla metodológica: afirmar únicamente lo trazable en la evidencia recibida.", Inches(1.0), Inches(6.37), Inches(10.7), Inches(0.25), 11, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 8)
    add_notes(slide, SLIDES[7][1])


def slide_tests(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 9, "Pruebas y escenarios documentados", "Informe del 4 de agosto de 2026")
    add_card(
        slide,
        Inches(0.72),
        Inches(1.88),
        Inches(5.7),
        Inches(2.72),
        "UAT-001 · Calcular y ver ruta",
        "Esperado\nDatos del bus y mapa con paradas numeradas.\n\n"
        "Visible\nCálculo para un bus y cinco paradas; no aparece el mapa de detalle esperado.",
        BLUE,
        17,
        12,
    )
    add_card(
        slide,
        Inches(6.68),
        Inches(1.88),
        Inches(5.7),
        Inches(2.72),
        "UAT-002 · Calcular y editar paradas",
        "Esperado\nGestión de paradas con opciones de edición.\n\n"
        "Visible\nDetalle y mapa; no se muestran controles de edición ni una operación guardada.",
        CYAN,
        17,
        12,
    )
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(4.95), Inches(11.66), Inches(1.15), NAVY, NAVY)
    add_text(slide, "2 DE 2 DECLARADOS APROBADOS", Inches(1.02), Inches(5.19), Inches(4.2), Inches(0.34), 17, WHITE, True)
    add_text(slide, "Estado del informe ≠ verificación visual completa", Inches(5.06), Inches(5.2), Inches(6.75), Inches(0.32), 16, AMBER, True, PP_ALIGN.RIGHT)
    add_text(slide, "Ambos casos corresponden al perfil administrador.", Inches(0.92), Inches(6.37), Inches(5.4), Inches(0.28), 11, GRAY, True)
    add_text(slide, "No usar «100 %» como cobertura del sistema.", Inches(6.72), Inches(6.37), Inches(5.35), Inches(0.28), 11, RED, True, PP_ALIGN.RIGHT)
    add_footer(slide, 9, "Fuente: informe funcional v1.0.0 citado en la tesis.")
    add_notes(slide, SLIDES[8][1])


def slide_results(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 10, "Resultados y alcance real", "Matriz de evidencia")
    columns = [
        ("RESPALDADO", GREEN, ["Rutas y mapas", "Ubicación observable", "Asistencia en interfaz"]),
        ("PARCIAL", YELLOW, ["Edición de paradas", "Persistencia de asistencia"]),
        ("PENDIENTE", GRAY, ["Abordaje persistente", "Rendimiento y producción", "Aceptación formal"]),
    ]
    for i, (title, color, items) in enumerate(columns):
        x = 0.72 + i * 4.0
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.93), Inches(3.65), Inches(3.9), WHITE, LINE)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.93), Inches(3.65), Inches(0.72), color, color)
        add_text(slide, title, Inches(x + 0.2), Inches(2.14), Inches(3.25), Inches(0.3), 15, WHITE, True, PP_ALIGN.CENTER)
        for j, item in enumerate(items):
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.3), Inches(2.98 + j * 0.82), Inches(0.32), Inches(0.32), color, color)
            add_text(slide, item, Inches(x + 0.78), Inches(2.91 + j * 0.82), Inches(2.55), Inches(0.45), 13, INK, j == 0)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.14), Inches(11.65), Inches(0.62), LIGHT, LIGHT)
    add_text(slide, "Conclusión restringida a un ambiente controlado y a las pantallas observadas.", Inches(0.95), Inches(6.31), Inches(11.2), Inches(0.28), 12, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 10)
    add_notes(slide, SLIDES[9][1])


def slide_conclusion(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_header(slide, 11, "Evaluación de hipótesis y trabajo futuro", "Conclusión defendible")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(1.9), Inches(5.45), Inches(2.38), NAVY, NAVY)
    add_text(slide, "HIPÓTESIS", Inches(1.05), Inches(2.2), Inches(2.2), Inches(0.3), 12, CYAN, True)
    add_text(slide, "RESPALDADA", Inches(1.05), Inches(2.72), Inches(4.7), Inches(0.55), 28, WHITE, True)
    add_text(slide, "dentro del alcance funcional documentado", Inches(1.05), Inches(3.42), Inches(4.55), Inches(0.36), 14, RGBColor(203, 218, 230), True)
    add_card(slide, Inches(6.48), Inches(1.9), Inches(5.9), Inches(2.38), "NO CERTIFICA", "Producción · reducción de tiempos · satisfacción · seguridad física · abordaje persistente", AMBER, 16, 14)
    add_text(slide, "Trabajo futuro", Inches(0.72), Inches(4.78), Inches(2.4), Inches(0.35), 17, NAVY, True)
    future = [
        ("01", "Verificar persistencia de asistencia y abordaje"),
        ("02", "Ampliar pruebas con padres, pilotos y monitores"),
        ("03", "Medir latencia, carga, estabilidad y autorizaciones"),
    ]
    for i, (num, text) in enumerate(future):
        x = 0.72 + i * 4.0
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.38), Inches(3.65), Inches(0.98), LIGHT, LINE)
        add_text(slide, num, Inches(x + 0.2), Inches(5.65), Inches(0.5), Inches(0.3), 12, CYAN, True)
        add_text(slide, text, Inches(x + 0.72), Inches(5.55), Inches(2.65), Inches(0.5), 12, INK, True)
    add_footer(slide, 11)
    add_notes(slide, SLIDES[10][1])


def slide_close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    add_text(slide, "APORTE VERIFICABLE", Inches(0.8), Inches(0.65), Inches(3.8), Inches(0.4), 13, CYAN, True)
    add_text(slide, "Una plataforma integrada\ncon evidencia trazable.", Inches(0.8), Inches(1.35), Inches(7.9), Inches(1.55), 34, WHITE, True)
    summary = [
        ("Observable", "Rutas · mapas · geolocalización · asistencia en interfaz", GREEN),
        ("Pendiente", "Abordaje persistente · rendimiento · aceptación integral", AMBER),
        ("Complementario", "Pagos fuera de la hipótesis evaluada", CYAN),
    ]
    for i, (title, body, color) in enumerate(summary):
        y = 3.35 + i * 0.9
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(1.72), Inches(0.5), color, color)
        add_text(slide, title.upper(), Inches(0.8), Inches(y + 0.11), Inches(1.72), Inches(0.25), 10, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(2.8), Inches(y + 0.03), Inches(7.8), Inches(0.45), 14, RGBColor(218, 229, 239), False)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(9.55), Inches(1.1), Inches(2.8), Inches(4.95), RGBColor(19, 47, 76), RGBColor(46, 78, 106))
    add_text(slide, "¿PREGUNTAS?", Inches(9.78), Inches(2.25), Inches(2.35), Inches(0.5), 22, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "Gracias", Inches(9.78), Inches(3.15), Inches(2.35), Inches(0.5), 20, CYAN, True, PP_ALIGN.CENTER)
    add_text(slide, "Transportes Génesis", Inches(9.78), Inches(4.18), Inches(2.35), Inches(0.35), 11, RGBColor(182, 201, 217), True, PP_ALIGN.CENTER)
    add_text(slide, "Universidad Galileo · 2026", Inches(0.8), Inches(6.75), Inches(4.3), Inches(0.3), 10, RGBColor(153, 176, 195))
    add_notes(slide, SLIDES[11][1])


def write_script() -> None:
    sections = [
        "# Guion de presentación de tesis",
        "",
        "Duración sugerida: 8–10 minutos, más preguntas.",
        "",
        "> Antes de exponer: completar facultad/escuela, carnés y asesor en la portada.",
        "",
    ]
    for index, (title, notes) in enumerate(SLIDES, start=1):
        sections.extend(
            [
                f"## Diapositiva {index}. {title}",
                "",
                notes,
                "",
                "**Idea de cierre:** " + [
                    "La exposición se limitará a los resultados documentados.",
                    "La necesidad es centralizar información operativa dispersa.",
                    "La hipótesis evalúa integración funcional, no impacto estadístico.",
                    "La arquitectura conecta interfaces, servicios, datos y mapas.",
                    "Cada perfil participa en un flujo operativo diferenciado.",
                    "La ubicación es observable con dependencias técnicas explícitas.",
                    "Asistencia observable no equivale a abordaje persistido.",
                    "La metodología separa evidencia, interpretación y límites.",
                    "El informe declara aprobación; las capturas son parciales.",
                    "El alcance real se comunica sin usar métricas no demostradas.",
                    "La hipótesis se respalda de forma acotada y verificable.",
                    "El principal aporte es integración con trazabilidad documental.",
                ][index - 1],
                "",
            ]
        )
    sections.extend(
        [
            "## Respuestas breves ante preguntas previsibles",
            "",
            "- **¿La UAT demuestra el sistema completo?** No. El informe documenta dos casos del administrador y declara ambos aprobados; las capturas ofrecen respaldo parcial.",
            "- **¿Por qué no hubo encuestas?** La investigación valida funcionamiento observable con la evidencia disponible; no afirma satisfacción ni aceptación general.",
            "- **¿El abordaje quedó validado?** No de forma persistente. Existe diseño e interfaz, pero falta una prueba que guarde y recupere el registro.",
            "- **¿Es tiempo real?** Se presenta como tiempo casi real porque existe latencia y depende de dispositivo, permisos y conectividad.",
            "- **¿La ruta es óptima?** El orden usa una heurística de vecino más cercano; es práctico, pero no garantiza el óptimo global.",
            "- **¿Los pagos confirman la hipótesis?** No. Son un módulo complementario y no fueron cubiertos por el informe funcional.",
            "",
        ]
    )
    SCRIPT_OUT.write_text("\n".join(sections), encoding="utf-8")


def export_pdf() -> bool:
    escaped_in = str(OUT).replace("'", "''")
    escaped_out = str(PDF_OUT).replace("'", "''")
    command = (
        "$ppt = New-Object -ComObject PowerPoint.Application; "
        "$ppt.Visible = -1; "
        f"$deck = $ppt.Presentations.Open('{escaped_in}', $true, $false, $false); "
        f"$deck.SaveAs('{escaped_out}', 32); "
        "$deck.Close(); $ppt.Quit(); "
        "[System.Runtime.Interopservices.Marshal]::ReleaseComObject($deck) | Out-Null; "
        "[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null"
    )
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-Command", command],
            check=True,
            timeout=90,
            capture_output=True,
            text=True,
        )
        return PDF_OUT.exists() and PDF_OUT.stat().st_size > 0
    except (subprocess.SubprocessError, OSError):
        return False


def main() -> None:
    for path in (ARCH, ATTENDANCE):
        if not path.exists():
            raise FileNotFoundError(path)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Defensa de tesis — Transportes Génesis"
    prs.core_properties.subject = "Presentación académica de resultados y alcance funcional"
    prs.core_properties.author = "Jonathan Samuel Villeda Pérez y José David Florián Secaida"
    prs.core_properties.keywords = "tesis, Transportes Génesis, geolocalización, asistencia, UAT"

    slide_portada(prs)
    slide_problem(prs)
    slide_hypothesis(prs)
    slide_architecture(prs)
    slide_roles(prs)
    slide_geo(prs)
    slide_attendance(prs)
    slide_method(prs)
    slide_tests(prs)
    slide_results(prs)
    slide_conclusion(prs)
    slide_close(prs)

    prs.save(OUT)
    write_script()
    pdf_ok = export_pdf()

    print(f"PowerPoint: {OUT}")
    print(f"Guion: {SCRIPT_OUT}")
    print(f"PDF: {PDF_OUT if pdf_ok else 'no exportado'}")
    print(f"Diapositivas: {len(prs.slides)}")


if __name__ == "__main__":
    main()
