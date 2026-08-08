"""Auditoría automática de los entregables finales de la tesis."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from pypdf import PdfReader


BASE = Path(__file__).resolve().parent
MD = BASE / "TESIS_TRANSPORTES_GENESIS_FINAL_11_AGOSTO.md"
DOCX = BASE / "TESIS_TRANSPORTES_GENESIS_FINAL_11_AGOSTO.docx"
PDF = BASE / "TESIS_TRANSPORTES_GENESIS_FINAL_11_AGOSTO.pdf"


def close(actual: float, expected: float, tolerance: float = 0.03) -> bool:
    return abs(actual - expected) <= tolerance


def main() -> None:
    for path in (MD, DOCX, PDF):
        assert path.exists() and path.stat().st_size > 0, f"Falta {path.name}"

    source = MD.read_text(encoding="utf-8")
    forbidden = (
        "[PEGAR",
        "UAT-__",
        "DEF-001",
        "Admin123",
        "Password:",
        "video_uso_geolocalizacion_1.mp4)",
        "video_geolocalizacion_2.mp4)",
        "video_pagos_1.mp4)",
    )
    for value in forbidden:
        assert value not in source, f"Contenido no permitido: {value}"

    required = (
        "Hipótesis de investigación (H1)",
        "Objetivo general",
        "Objetivos específicos",
        "la hipótesis de investigación queda respaldada dentro del alcance funcional",
        "# BIBLIOGRAFÍA",
        "# DATOS ADMINISTRATIVOS PENDIENTES",
        "# ACTA DE ACEPTACIÓN DE USUARIOS",
    )
    source_lower = source.lower()
    for value in required:
        assert value.lower() in source_lower, f"Falta contenido: {value}"

    citations = (
        "International Organization for Standardization",
        "Leaflet",
        "Microsoft",
        "OpenStreetMap contributors",
        "OWASP Foundation",
        "PCI Security Standards Council",
        "Project OSRM",
        "Universidad Galileo, FISICC",
        "World Wide Web Consortium",
        "Equipo QA y validadores técnicos regionales",
    )
    for citation in citations:
        assert source.count(citation) >= 2, f"Referencia sin cita: {citation}"

    document = Document(DOCX)
    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size and close(normal.font.size.pt, 12)
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE
    assert normal.paragraph_format.first_line_indent
    assert close(normal.paragraph_format.first_line_indent.cm, 0.63)
    assert normal.paragraph_format.alignment in (None, WD_ALIGN_PARAGRAPH.LEFT)
    assert "Tesis Cuadro" in document.styles
    assert "Tesis Figura" in document.styles

    for section in document.sections:
        assert close(section.top_margin.cm, 4)
        assert close(section.left_margin.cm, 4)
        assert close(section.bottom_margin.cm, 2.5)
        assert close(section.right_margin.cm, 2.5)
    assert all(
        section.start_type == WD_SECTION.ODD_PAGE
        for section in document.sections[2:]
    )

    page_number = document.sections[1]._sectPr.find(qn("w:pgNumType"))
    assert page_number is not None
    assert page_number.get(qn("w:fmt")) == "lowerRoman"
    assert page_number.get(qn("w:start")) == "8"
    for section in document.sections[3:]:
        continued_number = section._sectPr.find(qn("w:pgNumType"))
        assert (
            continued_number is None
            or qn("w:start") not in continued_number.attrib
        ), "La numeración se reinicia en un capítulo o apéndice"

    table_paragraphs = [
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    assert table_paragraphs
    assert all(
        paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        for paragraph in table_paragraphs
    )
    assert all(
        not paragraph.paragraph_format.first_line_indent
        or paragraph.paragraph_format.first_line_indent.twips == 0
        for paragraph in table_paragraphs
    )
    assert sum(
        len(paragraph._p.xpath(".//w:tab")) for paragraph in table_paragraphs
    ) == 0

    cover = document.paragraphs[:11]
    assert cover and all(
        paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER for paragraph in cover
    )
    assert len(document.inline_shapes) == 4
    assert not any("[IMAGEN:" in paragraph.text for paragraph in document.paragraphs)
    assert not any(
        "Actualice este campo" in paragraph.text for paragraph in document.paragraphs
    )
    assert not any(
        "No se encontraron entradas" in paragraph.text
        for paragraph in document.paragraphs
    )
    captions = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name in ("Tesis Cuadro", "Tesis Figura")
    ]
    assert captions
    assert len(captions) == 24
    index_entries = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Normal"
        and paragraph.text.startswith(("Cuadro ", "Figura "))
    ]
    assert len(index_entries) == 24
    page_references = [
        node
        for node in document.part._element.xpath(".//w:instrText")
        if "PAGEREF Cuadro_" in (node.text or "")
        or "PAGEREF Figura_" in (node.text or "")
    ]
    assert len(page_references) == 24

    links = [
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.is_external and "hyperlink" in relationship.reltype
    ]
    assert links
    assert all("Documento%20de%20Pruebas%20UAT.pdf" not in link for link in links)
    assert all(
        "video_uso_geolocalizacion_1.mp4" not in link
        and "video_geolocalizacion_2.mp4" not in link
        and "video_pagos_1.mp4" not in link
        for link in links
    )

    pdf = PdfReader(PDF)
    assert len(pdf.pages) >= 50
    first_page = pdf.pages[0].extract_text() or ""
    assert "UNIVERSIDAD GALILEO" in first_page
    assert not first_page.lstrip().startswith("CARÁTULA")
    pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "No se encontraron entradas" not in pdf_text
    assert pdf_text.count("Cuadro A.1") >= 2
    assert "CAPÍTULO 2" in pdf_text and "MARCO TEÓRICO" in pdf_text
    assert "ACTA DE ACEPTACIÓN DE USUARIOS" in pdf_text

    outline = pdf.outline
    assert outline, "El PDF no contiene marcadores de navegación"
    outline_titles = []

    def collect_outline_titles(items) -> None:
        for item in items:
            if isinstance(item, list):
                collect_outline_titles(item)
            else:
                title = getattr(item, "title", "")
                if title:
                    outline_titles.append(title)

    collect_outline_titles(outline)
    assert any("CAPÍTULO 1" in title for title in outline_titles)
    assert any("BIBLIOGRAFÍA" in title for title in outline_titles)

    internal_links = 0
    for page in pdf.pages:
        for reference in page.get("/Annots") or []:
            annotation = reference.get_object()
            if annotation.get("/Subtype") != "/Link":
                continue
            action = annotation.get("/A")
            if annotation.get("/Dest") is not None or (
                action is not None and action.get("/S") == "/GoTo"
            ):
                internal_links += 1
    assert internal_links >= 10, "El PDF no contiene suficientes enlaces internos"

    print("VALIDACIÓN FINAL APROBADA")
    print(f"Markdown: {MD.stat().st_size} bytes")
    print(
        f"Word: {DOCX.stat().st_size} bytes, "
        f"{len(document.sections)} secciones, "
        f"{len(document.tables)} tablas, "
        f"{len(document.inline_shapes)} figuras"
    )
    print(f"PDF: {PDF.stat().st_size} bytes")
    print(
        f"Navegación PDF: {len(outline_titles)} marcadores, "
        f"{internal_links} enlaces internos"
    )
    print(f"Hipervínculos externos: {len(set(links))}")


if __name__ == "__main__":
    main()
