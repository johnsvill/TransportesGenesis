"""
Genera versiones Word con formato formal (márgenes, Times New Roman 12, interlineado 1.5,
texto justificado, portada desde metadatos YAML) a partir de los Markdown de la tesis.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent

MARGIN_TOP = Cm(2.5)
MARGIN_BOTTOM = Cm(2.5)
MARGIN_LEFT = Cm(3.0)
MARGIN_RIGHT = Cm(2.5)

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
HEADING1_SIZE = Pt(14)
HEADING2_SIZE = Pt(13)
HEADING3_SIZE = Pt(12)


def parse_yaml_frontmatter(content: str) -> tuple[dict[str, str], str]:
    meta: dict[str, str] = {}
    if not content.startswith("---"):
        return meta, content
    end = content.find("---", 3)
    if end == -1:
        return meta, content
    block = content[3:end].strip()
    body = content[end + 3 :].lstrip("\n")
    for line in block.splitlines():
        if ":" in line and not line.strip().startswith("formato"):
            key, val = line.split(":", 1)
            meta[key.strip()] = val.strip().strip('"')
    return meta, body


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = MARGIN_TOP
    section.bottom_margin = MARGIN_BOTTOM
    section.left_margin = MARGIN_LEFT
    section.right_margin = MARGIN_RIGHT
    add_page_number_footer(section)


def add_page_number_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=Pt(10))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run._r.append(fld)


def set_run_font(run, bold=False, italic=False, size=None, color=None):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = size or FONT_SIZE
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def parse_inline(text: str, paragraph):
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|[^*]+)")
    for part in pattern.findall(text):
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, italic=True)
        elif part:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_body_paragraph(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0.63)
    parse_inline(text.strip(), p)
    return p


def add_heading(doc, text: str, level: int, page_break: bool = False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(12 if level <= 2 else 6)
    p.paragraph_format.space_after = Pt(6)
    sizes = {1: HEADING1_SIZE, 2: HEADING2_SIZE, 3: HEADING3_SIZE}
    run = p.add_run(text.strip())
    set_run_font(run, bold=True, size=sizes.get(level, FONT_SIZE))
    return p


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            parse_inline(cell_text.strip(), p)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def add_cover_page(doc, title: str, subtitle: str, doc_type: str, meta: dict[str, str]):
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TRANSPORTES GÉNESIS")
    set_run_font(run, bold=True, size=Pt(16))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, bold=True, size=Pt(14))

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subtitle)
        set_run_font(run, size=Pt(12))

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(doc_type)
    set_run_font(run, bold=True, size=Pt(12))

    for _ in range(6):
        doc.add_paragraph()

    fields = [
        ("Institución:", meta.get("institucion", "[Nombre de la universidad]")),
        ("Carrera:", meta.get("carrera", "[Nombre de la carrera]")),
        ("Autor(es):", meta.get("autor", "[Nombre del autor]")),
        ("Asesor:", meta.get("asesor", "[Nombre del asesor]")),
        ("Lugar:", meta.get("lugar", "San Cristóbal, departamento de San Marcos, Guatemala")),
        ("Fecha:", meta.get("fecha", "Julio 2026")),
    ]
    for label, value in fields:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        r1 = p.add_run(f"{label} ")
        set_run_font(r1, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2)
        if "[" in value:
            print(f"  AVISO: completar placeholder en portada: {label} {value}")

    doc.add_page_break()


def markdown_to_docx(md_path: Path, docx_path: Path, cover: dict):
    raw = md_path.read_text(encoding="utf-8")
    meta, content = parse_yaml_frontmatter(raw)

    doc = Document()
    setup_document(doc)
    add_cover_page(
        doc,
        cover["title"],
        cover.get("subtitle", ""),
        cover["doc_type"],
        meta,
    )

    lines = content.splitlines()
    i = 0
    table_rows: list[list[str]] = []
    in_table = False
    list_buffer: list[str] = []

    def flush_list():
        nonlocal list_buffer
        for item in list_buffer:
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            parse_inline(item, p)
        list_buffer = []

    while i < len(lines):
        line = lines[i]

        if in_table:
            if line.strip().startswith("|") and not re.match(r"^\|\s*[-:| ]+\|\s*$", line.strip()):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                table_rows.append(cells)
                i += 1
                continue
            add_table(doc, table_rows)
            table_rows = []
            in_table = False
            continue

        if line.strip().startswith("|") and not re.match(r"^\|\s*[-:| ]+\|\s*$", line.strip()):
            in_table = True
            table_rows = [[c.strip() for c in line.strip().strip("|").split("|")]]
            i += 1
            if i < len(lines) and re.match(r"^\|\s*[-:| ]+\|\s*$", lines[i].strip()):
                i += 1
            continue

        if line.strip().startswith("- ") or line.strip().startswith("* "):
            if not line.strip().startswith("*Documento"):
                list_buffer.append(line.strip()[2:])
            i += 1
            continue
        elif list_buffer:
            flush_list()

        if re.match(r"^\d+\.\s", line.strip()):
            p = doc.add_paragraph(style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            parse_inline(re.sub(r"^\d+\.\s", "", line.strip()), p)
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            text = line[2:].strip()
            is_chapter = text.lower().startswith("capítulo")
            add_heading(doc, text, 1, page_break=is_chapter)
            i += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 2)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
            i += 1
            continue

        if line.strip().startswith(">"):
            text = line.strip().lstrip(">").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            run = p.add_run(text)
            set_run_font(run, italic=True)
            i += 1
            continue

        if line.strip().startswith("*Documento"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.strip().strip("*"))
            set_run_font(run, italic=True, size=Pt(11))
            i += 1
            continue

        if line.strip():
            add_body_paragraph(doc, line.strip())
        i += 1

    if list_buffer:
        flush_list()
    if table_rows:
        add_table(doc, table_rows)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"Generado: {docx_path}")


def main():
    docs = [
        {
            "md": BASE / "01_Introduccion_y_Marco_Teorico.md",
            "docx": BASE / "01_Introduccion_y_Marco_Teorico.docx",
            "cover": {
                "title": "Introducción y Marco Teórico",
                "subtitle": "Sistema web integrado de geolocalización y gestión de pagos para transporte escolar",
                "doc_type": "Documento de Tesis — Capítulos 1 y 2",
            },
        },
        {
            "md": BASE / "pruebas_tesis" / "PLAN_PRUEBAS_UAT_CARGA_RENDIMIENTO.md",
            "docx": BASE / "pruebas_tesis" / "PLAN_PRUEBAS_UAT_CARGA_RENDIMIENTO.docx",
            "cover": {
                "title": "Plan de Pruebas: UAT, Carga y Rendimiento",
                "subtitle": "Validación del sistema — Segmento: padres con hijos en primer grado",
                "doc_type": "Anexo de Tesis — Plan de Pruebas",
            },
        },
    ]

    for item in docs:
        if not item["md"].exists():
            print(f"ERROR: no existe {item['md']}")
            continue
        markdown_to_docx(item["md"], item["docx"], item["cover"])


if __name__ == "__main__":
    main()
