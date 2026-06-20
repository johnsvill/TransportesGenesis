"""Convierte Proyecto_Transportes_Genesis_v3_ACTUALIZADO.md a .docx"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

MD_PATH = Path(__file__).parent / "Proyecto_Transportes_Genesis_v3_ACTUALIZADO.md"
DOCX_PATH = Path(__file__).parent / "Proyecto_Transportes_Genesis_v3_ACTUALIZADO.docx"


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for level, size in [(1, 18), (2, 14), (3, 12), (4, 11)]:
        name = f"Heading {level}"
        if name in doc.styles:
            doc.styles[name].font.name = "Calibri"
            doc.styles[name].font.size = Pt(size)
            doc.styles[name].font.bold = True


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.rows[i].cells[j].text = cell.strip()
    doc.add_paragraph()


def parse_md_lines(lines):
    i = 0
    in_code = False
    code_lines = []
    code_lang = ""
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line.strip()[3:].strip()
                code_lines = []
            else:
                yield ("code", code_lang, "\n".join(code_lines))
                in_code = False
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("# "):
            yield ("h1", line[2:].strip())
        elif line.startswith("## "):
            yield ("h2", line[3:].strip())
        elif line.startswith("### "):
            yield ("h3", line[4:].strip())
        elif line.startswith("#### "):
            yield ("h4", line[5:].strip())
        elif line.strip() == "---":
            yield ("hr",)
        elif line.startswith("|") and "|" in line[1:]:
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                if re.match(r"^\|[-:\s|]+\|$", row):
                    i += 1
                    continue
                cells = [c.strip() for c in row.strip("|").split("|")]
                table_rows.append(cells)
                i += 1
            yield ("table", table_rows)
            continue
        elif line.strip().startswith("- [ ]") or line.strip().startswith("- [x]"):
            yield ("li", line.strip()[2:].strip())
        elif line.strip().startswith("- "):
            yield ("li", line.strip()[2:].strip())
        elif line.strip().startswith("> "):
            yield ("quote", line.strip()[2:].strip())
        elif line.strip():
            yield ("p", line.strip())
        i += 1


def add_rich_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)
    return p


def convert():
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    setup_styles(doc)

    # Portada breve
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = title.add_run("Transportes Génesis\nEspecificación del Proyecto v3.0")
    r.bold = True
    r.font.size = Pt(22)
    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sub.add_run("Alineado al código implementado y presentación Gestión de Pagos — 20/06/2026")
    doc.add_page_break()

    for item in parse_md_lines(lines):
        kind = item[0]
        if kind == "h1":
            doc.add_heading(item[1], level=1)
        elif kind == "h2":
            doc.add_heading(item[1], level=2)
        elif kind == "h3":
            doc.add_heading(item[1], level=3)
        elif kind == "h4":
            doc.add_heading(item[1], level=4)
        elif kind == "p":
            add_rich_paragraph(doc, item[1])
        elif kind == "li":
            add_rich_paragraph(doc, item[1], style="List Bullet")
        elif kind == "quote":
            p = add_rich_paragraph(doc, item[1])
            for run in p.runs:
                run.italic = True
        elif kind == "table":
            add_table(doc, item[1])
        elif kind == "code":
            lang, body = item[1], item[2]
            if lang in ("mermaid", "") and ("flowchart" in body or "sequenceDiagram" in body or "erDiagram" in body or "xychart" in body or "pie" in body):
                note = doc.add_paragraph()
                note.add_run("[Diagrama — ver versión Markdown o recrear en draw.io / PowerPoint]").italic = True
                p = doc.add_paragraph(body)
                p.style = "No Spacing"
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            else:
                p = doc.add_paragraph(body)
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
        elif kind == "hr":
            doc.add_paragraph()

    doc.save(str(DOCX_PATH))
    print(f"Created: {DOCX_PATH} ({DOCX_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    convert()
